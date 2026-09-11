from pathlib import Path
import json
import pandas as pd
import pytest
from lancet_funding.dimensions import Client
from lancet_funding.workflow import run_all,extract_stage,analyse_stage
from lancet_funding.io import read_table,InputError,write_csv
from lancet_funding.comparison import compare_master
from lancet_funding.excel import export_xlsx


def test_end_to_end_from_mock_api_and_resume(inputs):
    cfg,api,run=inputs
    c=Client(run/'raw'/'cache',dsl=api,pause=0)
    manifest=run_all(cfg,run,client=c,figures=False,data_label='synthetic')
    assert manifest['status']=='complete_analysis_not_historical_validation'
    assert manifest['historical_report_reproduction_certified'] is False
    files=list((run/'results').rglob('*.xlsx'))
    assert len(files)>=30
    main=read_table(run/'results/tables/main_global.csv')
    assert set(main.scenario_id.astype(str))=={'14'}
    f=main[(main.kind=='funder')&(main.year==2016)].iloc[0]
    r=main[(main.kind=='recipient')&(main.year==2016)].iloc[0]
    # Explicit arithmetic, independent of the aggregator implementation.
    assert f.climate_count==7 # six retained unique grants, one credited to two funders
    assert r.climate_count==5 # zero/unknown funding removed; global grants counted once
    assert f.climate_funding_usd==pytest.approx((100+50+30+20+45*2)*1.06)
    assert r.climate_funding_usd==pytest.approx((100+50+30+20+45)*1.06)
    n=len(api.calls)
    extract_stage(cfg,run,client=c)
    assert len(api.calls)==n
    with pytest.raises(InputError):extract_stage({**cfg,'start_year':2000},run,client=c)
    baseline=run/'results/funder/CLEANED_RAW_FUNDER_DATASET.xlsx'
    current=run/'results/processed/funder_master.csv'
    result=compare_master(baseline,current,'funder',run/'comparison')
    assert result['status']=='PASS'
    changed=read_table(current);changed.loc[changed.index[changed.scenario_id.astype(str).eq('14')][0],'health_count']+=1
    write_csv(changed,run/'changed.csv')
    assert compare_master(baseline,run/'changed.csv','funder',run/'comparison')['status']=='DIFFERENCES_FOUND'


def test_extraction_without_reference_assets(inputs):
    cfg,api,run=inputs
    Path(cfg['manual_exclusions']).unlink()
    c=Client(run/'raw'/'cache',dsl=api,pause=0)
    extract_stage(cfg,run,client=c)
    assert (run/'raw/COMPLETE.json').exists()
    with pytest.raises(InputError):analyse_stage(cfg,run,figures=False)


def test_xlsx_untrusted_strings_not_formulas(tmp_path):
    from openpyxl import load_workbook
    p=tmp_path/'check.xlsx'
    export_xlsx({'data':pd.DataFrame({'title':['=1+1','@SUM(A1:A2)'],'count':[1,2]})},p)
    wb=load_workbook(p);assert wb['data']['A2'].data_type=='s'
    assert wb['data']['A2'].value=='=1+1'
    assert wb['data'].freeze_panes=='A2'


def test_cli_run_from_zero_through_mock_http(inputs,monkeypatch,tmp_path):
    import requests
    from lancet_funding.cli import main
    cfg,api,run=inputs
    cfg['analysis_start']=2010;cfg['pause_seconds']=0
    path=tmp_path/'config.json';path.write_text(json.dumps(cfg))
    class Response:
        status_code=200
        headers={}
        def __init__(self,obj):self.obj=obj
        def json(self):return self.obj
    class Session:
        def post(self,url,**kwargs):
            if url.endswith('/api/auth'):
                assert kwargs['json']=={'key':'synthetic-cli-key'}
                return Response({'token':'synthetic-cli-token'})
            assert url.endswith('/api/dsl/v2')
            assert kwargs['headers']['Authorization']=='JWT synthetic-cli-token'
            return Response(api.query(kwargs['data'].decode()).json)
    monkeypatch.setattr(requests,'Session',Session)
    monkeypatch.setenv('DIMENSIONS_API_KEY','synthetic-cli-key')
    monkeypatch.setenv('DIMENSIONS_ENDPOINT','https://app.dimensions.ai')
    assert main(['run','--config',str(path),'--run-dir',str(run),'--no-figures'])==0
    assert (run/'results/ALL_TABLES.xlsx').exists()
    for p in run.rglob('*'):
        if p.is_file() and p.suffix in {'.json','.jsonl','.csv'}:
            text=p.read_text()
            assert 'synthetic-cli-key' not in text and 'synthetic-cli-token' not in text


def test_long_band_share_cells_are_formatted_as_percent(tmp_path):
    from openpyxl import load_workbook
    p=tmp_path/'band.xlsx'
    export_xlsx({'data':pd.DataFrame({'metric':['funding_share','climate_count'],'band_min':[.0013,15]})},p)
    wb=load_workbook(p)
    assert wb['data']['B2'].value==.0013 and wb['data']['B2'].number_format=='0.000%'
    assert wb['data']['B3'].number_format.startswith('#,##0;')
