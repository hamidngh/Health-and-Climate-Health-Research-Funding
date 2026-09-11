"""Portable XLSX exports. Calculations are in Python; CSV is the canonical table.

The grant-level CSV/JSONL retain full strings; Excel cells have a 32767-character
limit. Literal strings are typed as text so untrusted grant titles cannot become
Excel formulas. No macros, external links or hidden sheets are created.
"""
from pathlib import Path
import json
import math
import re
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .io import InputError, write_csv, stamp


def export_xlsx(sheets: dict, path, *, status='Locally computed results; historical agreement not certified'):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    wb=Workbook(write_only=True)
    used=set()
    for name,df in sheets.items():
        name=re.sub(r'[\\/*?:\[\]]','_',name)[:24]
        chunks=max(1,math.ceil(len(df)/1_000_000))
        for k in range(chunks):
            title=name if chunks==1 else name+'_'+str(k+1)
            if title in used:raise InputError('Duplicate Excel sheet name: '+title)
            used.add(title);ws=wb.create_sheet(title)
            ws.freeze_panes='A2';ws.row_dimensions[1].height=32
            header=[]
            for i,col in enumerate(df.columns,1):
                cell=WriteOnlyCell(ws,str(col));cell.font=Font(bold=True,color='FFFFFF')
                cell.fill=PatternFill('solid',fgColor='214E64');cell.alignment=Alignment(wrap_text=True)
                header.append(cell);ws.column_dimensions[get_column_letter(i)].width=min(38,max(14,len(str(col))+2))
            ws.append(header)
            part=df.iloc[k*1_000_000:(k+1)*1_000_000]
            metric_index=list(df.columns).index('metric') if 'metric' in df else None
            for values in part.itertuples(index=False,name=None):
                row=[]
                metric=str(values[metric_index]) if metric_index is not None else ''
                long_value_columns={'value','main','band_min','band_max','alternative_min','alternative_max',
                                    'all_11_min','all_11_max','baseline','current','difference'}
                for col,value in zip(df.columns,values):
                    if isinstance(value,(list,dict,tuple)):value=json.dumps(value,ensure_ascii=False)
                    if isinstance(value,np.generic):value=value.item()
                    if pd.isna(value):value=None
                    if isinstance(value,str):value=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',value)[:32767]
                    cell=WriteOnlyCell(ws,value)
                    if isinstance(value,str):cell.data_type='s'
                    elif isinstance(value,(int,float)) and not isinstance(value,bool):
                        if col in long_value_columns and metric.endswith('_share'):
                            cell.number_format='0.000%'
                        elif col in long_value_columns and metric.endswith('_count'):
                            cell.number_format='#,##0;[Red](#,##0);0'
                        elif (str(col).endswith('_share') or str(col).endswith('_share_legacy_zero')) or col in ['Ratio_Count','Ratio_Funding']:cell.number_format='0.00%'
                        elif 'year' in str(col).lower():cell.number_format='0'
                        elif 'count' in str(col).lower():cell.number_format='#,##0;[Red](#,##0);0'
                        else:cell.number_format='#,##0.00;[Red](#,##0.00);0'
                    row.append(cell)
                ws.append(row)
            if len(df.columns):ws.auto_filter.ref=f'A1:{get_column_letter(len(df.columns))}{len(part)+1}'
    meta=wb.create_sheet('README');meta.column_dimensions['A'].width=24;meta.column_dimensions['B'].width=100
    for row in [
        ['Indicator','Main approach (14) and ten sensitivity approaches'],['Generated UTC',stamp()],['Status',status],
        ['Funding','USD award values assigned to grant start year, not annual expenditure'],
        ['Ratios','Shares are fractions; *_pct fields are percentage values. Undefined ratios are blank.'],
        ['Data source','https://docs.dimensions.ai/dsl/datasource-grants.html'],
        ['Full text','Excel strings are truncated at 32767 characters; use CSV/JSONL for full raw text.'],
        ['Calculation','All results are computed by the same Python code used for CSV export.']]:meta.append(row)
    tmp=path.with_name(path.stem+'.tmp.xlsx');wb.save(tmp);tmp.replace(path)


def export_pair(df, stem, *, status='Locally computed results; historical agreement not certified'):
    stem=Path(stem)
    serial=df.copy()
    for col in serial:
        serial[col]=serial[col].map(lambda x:json.dumps(x,ensure_ascii=False) if isinstance(x,(list,dict,tuple)) else x)
    write_csv(serial,stem.with_suffix('.csv'))
    export_xlsx({'data':df},stem.with_suffix('.xlsx'),status=status)
